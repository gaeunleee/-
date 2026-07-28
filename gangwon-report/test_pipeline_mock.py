#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
네트워크 접속 없이(가짜 API 응답으로) fetch_and_build_report.py의 필터링/엑셀 생성 로직을 검증한다.
실제 나라장터 데이터가 아닌 테스트용 샘플 데이터이며, 결과는 output/test_mock_output.xlsx 로 저장된다
(배포용 산출물인 강원도_전시체험_협상계약_2025_2026.xlsx 와는 별개 파일).
"""
from pathlib import Path

from openpyxl import load_workbook

import fetch_and_build_report as m

SAMPLE_ITEMS = [
    # 1) 지역+키워드+업종코드 모두 매칭 -> AND/OR 둘 다 통과해야 함
    {
        "bidNtceNo": "R25AA00000001",
        "bidNtceNm": "춘천 미디어아트 체험관 조성 용역",
        "ntceInsttNm": "강원특별자치도청",
        "dmndInsttNm": "강원특별자치도청",
        "cntrctCnclsMthdNm": "협상에의한계약",
        "asignBdgtAmt": "350,000,000",
        "bidNtceDate": "20250310",
        "bidClseDate": "20250401",
        "prtcptPsblRgnNm": "강원특별자치도",
        "bidNtceDtlUrl": "https://www.g2b.go.kr/example/1",
        "bidprcPsblIndstrytyNm": "전시사업자, 소프트웨어사업자",
    },
    # 2) 지역만 매칭, 키워드/업종코드 없음 -> AND/OR 둘 다 탈락
    {
        "bidNtceNo": "R25AA00000002",
        "bidNtceNm": "강릉시 도로 포장 공사",
        "ntceInsttNm": "강릉시청",
        "dmndInsttNm": "강릉시청",
        "cntrctCnclsMthdNm": "최저가낙찰",
        "asignBdgtAmt": "1,200,000,000",
        "bidNtceDate": "20250315",
        "bidClseDate": "20250405",
        "prtcptPsblRgnNm": "강원특별자치도 강릉시",
        "bidNtceDtlUrl": "https://www.g2b.go.kr/example/2",
        "bidprcPsblIndstrytyNm": "토목공사업",
    },
    # 3) 지역+키워드만 매칭 (업종코드 없음) -> OR만 통과, AND는 탈락
    {
        "bidNtceNo": "R25AA00000003",
        "bidNtceNm": "원주 도서관 리모델링 설계 용역",
        "ntceInsttNm": "원주시청",
        "dmndInsttNm": "원주시청",
        "cntrctCnclsMthdNm": "협상에의한계약",
        "asignBdgtAmt": "80,000,000",
        "bidNtceDate": "20250320",
        "bidClseDate": "20250410",
        "prtcptPsblRgnNm": "강원특별자치도 원주시",
        "bidNtceDtlUrl": "https://www.g2b.go.kr/example/3",
        "bidprcPsblIndstrytyNm": "건축설계업",
    },
    # 4) 지역이 강원이 아님 (다른 조건은 다 맞음) -> 지역 불일치로 무조건 탈락
    {
        "bidNtceNo": "R25AA00000004",
        "bidNtceNm": "서울 전시관 조형물 제작설치 용역",
        "ntceInsttNm": "서울시청",
        "dmndInsttNm": "서울시청",
        "cntrctCnclsMthdNm": "협상에의한계약",
        "asignBdgtAmt": "200,000,000",
        "bidNtceDate": "20250325",
        "bidClseDate": "20250415",
        "prtcptPsblRgnNm": "서울특별시",
        "bidNtceDtlUrl": "https://www.g2b.go.kr/example/4",
        "bidprcPsblIndstrytyNm": "전시사업자",
    },
    # 5) 페이지 경계 등으로 중복 수집되었다고 가정 (1번과 동일 공고번호) -> dedupe로 1건만 남아야 함
    {
        "bidNtceNo": "R25AA00000001",
        "bidNtceNm": "춘천 미디어아트 체험관 조성 용역",
        "ntceInsttNm": "강원특별자치도청",
        "dmndInsttNm": "강원특별자치도청",
        "cntrctCnclsMthdNm": "협상에의한계약",
        "asignBdgtAmt": "350,000,000",
        "bidNtceDate": "20250310",
        "bidClseDate": "20250401",
        "prtcptPsblRgnNm": "강원특별자치도",
        "bidNtceDtlUrl": "https://www.g2b.go.kr/example/1",
        "bidprcPsblIndstrytyNm": "전시사업자, 소프트웨어사업자",
    },
]


def run(match_mode: str) -> list[m.FilteredRow]:
    m.MATCH_MODE = match_mode
    seen = set()
    deduped = []
    for item in SAMPLE_ITEMS:
        no = item["bidNtceNo"]
        if no in seen:
            continue
        seen.add(no)
        deduped.append(("용역", item))

    rows = []
    for bt, item in deduped:
        row = m.evaluate(bt, item)
        if row:
            rows.append(row)
    return rows


def main() -> int:
    ok = True

    rows_and = run("AND")
    print(f"[AND 모드] 통과 {len(rows_and)}건: {[r.notice_no for r in rows_and]}")
    if [r.notice_no for r in rows_and] != ["R25AA00000001"]:
        print("  FAIL: AND 모드는 1번 공고만 통과해야 함")
        ok = False
    else:
        print("  PASS")

    rows_or = run("OR")
    print(f"[OR 모드] 통과 {len(rows_or)}건: {[r.notice_no for r in rows_or]}")
    if sorted(r.notice_no for r in rows_or) != ["R25AA00000001", "R25AA00000003"]:
        print("  FAIL: OR 모드는 1번, 3번 공고가 통과해야 함")
        ok = False
    else:
        print("  PASS")

    # 예산 문자열 콤마 파싱, 엑셀 저장까지 엔드투엔드 확인
    out_path = m.OUTPUT_DIR / "test_mock_output.xlsx"
    m.write_excel(rows_and, out_path)
    if not out_path.exists():
        print("  FAIL: 엑셀 파일이 생성되지 않음")
        ok = False
    else:
        wb = load_workbook(out_path)
        ws = wb.active
        header = [c.value for c in ws[1]]
        expected_header = [h for h, _k, _w in m.COLUMNS]
        if header != expected_header:
            print(f"  FAIL: 헤더 불일치 {header} != {expected_header}")
            ok = False
        budget_cell = ws.cell(row=2, column=6)
        if budget_cell.value != 350_000_000:
            print(f"  FAIL: 예산이 정수로 변환되지 않음 (실제: {budget_cell.value!r})")
            ok = False
        else:
            print(f"  PASS: 엑셀 저장/헤더/예산 파싱 확인 완료 -> {out_path}")

    print("\n결과:", "PASS" if ok else "FAIL")
    return 0 if ok else 1


if __name__ == "__main__":
    raise SystemExit(main())
