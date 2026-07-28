#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
강원도 전시/체험 관련 나라장터 입찰공고(용역/공사) 수집 -> 필터링 -> 엑셀 저장

사용법:
    export NARA_SERVICE_KEY="발급받은 인증키(Decoding)"
    pip install -r requirements.txt
    python fetch_and_build_report.py

    # 이미 받아둔 원본 데이터로 필터/엑셀만 다시 만들고 싶을 때 (API 재호출 없음):
    python fetch_and_build_report.py --skip-fetch

    # 중간에 실패했을 때 이어서 하기 (기본 동작 - cache에 남은 진행상황을 자동 재사용)
    python fetch_and_build_report.py

주의:
    - 실제 응답 필드명(계약방법/지역제한 등)은 FIELD_CANDIDATES 에 "후보"로 정리해뒀습니다.
      최초 실행 시 첫 응답의 원본 필드 목록을 콘솔에 출력하니, 값이 이상하면 그 목록을 보고
      FIELD_CANDIDATES 를 실제 필드명으로 수정하세요.
    - API 호출 사이 REQUEST_INTERVAL_SEC 만큼 대기해 과다호출을 피합니다.
    - 15일 단위 창 x 2개 오퍼레이션 = 약 100회 내외 호출이 필요해 몇 분 정도 걸릴 수 있습니다.
"""
from __future__ import annotations

import argparse
import json
import os
import sys
import time
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any, Iterable

import requests
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# ---------------------------------------------------------------------------
# 설정
# ---------------------------------------------------------------------------

BASE_URL = "http://apis.data.go.kr/1230000/ad/BidPublicInfoService"
OPERATIONS: dict[str, str] = {
    "용역": "getBidPblancListInfoServcPPSSrch",
    "공사": "getBidPblancListInfoCnstwkPPSSrch",
}

START_DATE = date(2025, 1, 1)
END_DATE = date(2026, 12, 31)
WINDOW_DAYS = 15  # API 특성상 한 번에 15일치까지만 조회 가능

NUM_OF_ROWS = 999
MAX_PAGES = 10  # 창 하나(15일)에서 999건 x 10페이지 = 9990건 넘게 나오는 경우는 사실상 없음(안전장치)

REQUEST_INTERVAL_SEC = 0.4
TIMEOUT_SEC = 20
MAX_RETRIES = 4
RETRY_BASE_DELAY_SEC = 1.5
RETRYABLE_RESULT_CODES = {"04", "05", "22"}
SUCCESS_RESULT_CODES = {"00", "03"}  # 03 = 데이터 없음(정상)

REGION_KEYWORD = "강원"

TITLE_KEYWORDS = [
    "전시", "체험", "제작설치", "전시관", "체험관", "콘텐츠",
    "미디어아트", "조성", "디자인", "도서관", "조형물",
]

# 세부품명 코드: 참고용으로 정의는 해두지만, 이 스크립트가 조회하는 용역/공사 API 응답에는
# 세부품명번호(물품 전용 필드)가 원천적으로 존재하지 않아 실질적으로 매칭되지 않습니다.
# (물품까지 포함하려면 getBidPblancListInfoThngPPSSrch 오퍼레이션을 추가해야 합니다.)
PRODUCT_CODES = {
    "5512190301": "안내전광판", "3911160501": "LED경관조명기구", "4511161601": "비디오프로젝터",
    "4511189301": "영상정보디스플레이장치", "4924159701": "조합놀이대", "4924159801": "기타놀이기구",
    "5610150201": "소파", "5610150701": "책장", "5610170501": "유리진열장", "5612100201": "카운터",
    "5612101501": "이동식서가", "5612109901": "이동식서가", "5612190201": "전시용진열대",
    "6010393101": "동물표본", "6010440301": "화석표본", "6010989901": "실물모형및전시물",
    "6012100201": "조형물", "6014119501": "가상현실장치", "7215409902": "전시홍보관설치서비스",
    "8014198801": "전시회기획및대행서비스", "9015180201": "전시회기획및대행서비스", "8214150201": "디자인서비스",
}

INDUSTRY_CODES = {
    "0006": "실내건축공사업", "4990": "실내건축공사업",
    "1440": "금속창호·지붕건축물조립공사업", "4991": "금속창호·지붕건축물조립공사업",
    "1426": "소프트웨어사업자", "1468": "소프트웨어사업자", "1469": "소프트웨어사업자", "1470": "소프트웨어사업자",
    "6484": "공공디자인 전문회사", "6815": "전시사업자",
}
INDUSTRY_NAMES = sorted(set(INDUSTRY_CODES.values()))

# "아래 조건 모두 만족" 을 문자 그대로 지역 AND 키워드 AND 코드 로 해석한 기본값입니다.
# 결과가 너무 적으면 "OR" 로 바꿔서 재실행하세요 (재호출 없이 캐시로 바로 재필터링됩니다).
#   "AND": 지역 AND 키워드 AND 업종코드 모두 만족해야 채택
#   "OR" : 지역은 필수 + (키워드 또는 업종코드 중 하나만 만족해도 채택)
MATCH_MODE = os.environ.get("MATCH_MODE", "AND").upper()

OUTPUT_FILENAME = "강원도_전시체험_협상계약_2025_2026.xlsx"

SCRIPT_DIR = Path(__file__).resolve().parent
CACHE_DIR = SCRIPT_DIR / "cache"
OUTPUT_DIR = SCRIPT_DIR / "output"
RAW_CACHE_PATH = CACHE_DIR / "raw_items.jsonl"
PROGRESS_PATH = CACHE_DIR / "progress.json"

# data.go.kr 응답 필드명은 서비스 버전에 따라 조금씩 다를 수 있어 후보 목록으로 관리합니다.
# 실행 시 첫 응답의 실제 키 목록을 콘솔에 출력하므로, 값이 비어 보이면 여기에 실제 필드명을 추가하세요.
FIELD_CANDIDATES: dict[str, list[str]] = {
    "notice_no": ["bidNtceNo"],
    "title": ["bidNtceNm"],
    "org": ["ntceInsttNm"],
    "demand_org": ["dmndInsttNm"],
    "contract_method": ["cntrctCnclsMthdNm"],
    "budget": ["asignBdgtAmt", "presmptPrce", "bssamt", "bssAmt"],
    "notice_date": ["bidNtceDate", "bidNtceBgnDt", "bidNtceBgn"],
    "close_date": ["bidClseDate", "bidClseDt"],
    "region_limit": ["prtcptPsblRgnNm", "rgnLmtYn"],
    "detail_url": ["bidNtceDtlUrl", "bidNtceUrl"],
    "industry_text": ["bidprcPsblIndstrytyNm", "indstrytyLmtYn"],
}


# ---------------------------------------------------------------------------
# 유틸
# ---------------------------------------------------------------------------


def log(msg: str) -> None:
    print(f"[{datetime.now().strftime('%H:%M:%S')}] {msg}", flush=True)


def date_windows(start: date, end: date, window_days: int) -> list[tuple[date, date]]:
    windows = []
    cur = start
    while cur <= end:
        win_end = min(cur + timedelta(days=window_days - 1), end)
        windows.append((cur, win_end))
        cur = win_end + timedelta(days=1)
    return windows


def to_api_datetime(d: date, end_of_day: bool) -> str:
    return f"{d.strftime('%Y%m%d')}{'2359' if end_of_day else '0000'}"


def pick(item: dict[str, Any], candidates: list[str]) -> str | None:
    for key in candidates:
        v = item.get(key)
        if isinstance(v, str) and v.strip():
            return v.strip()
        if isinstance(v, (int, float)):
            return str(v)
    return None


# ---------------------------------------------------------------------------
# API 호출
# ---------------------------------------------------------------------------


class ApiResultError(Exception):
    def __init__(self, result_code: str, result_msg: str):
        super().__init__(f"resultCode={result_code} resultMsg={result_msg}")
        self.result_code = result_code


def normalize_items(items_field: Any) -> list[dict[str, Any]]:
    if items_field is None or items_field == "":
        return []
    if isinstance(items_field, list):
        return [i for i in items_field if isinstance(i, dict)]
    if isinstance(items_field, dict):
        if "item" in items_field:
            return normalize_items(items_field["item"])
        return [items_field]
    return []


def call_once(session: requests.Session, operation: str, service_key: str, params: dict[str, str]) -> dict[str, Any]:
    url = f"{BASE_URL}/{operation}"
    query = {"serviceKey": service_key, "type": "json", **params}
    resp = session.get(url, params=query, timeout=TIMEOUT_SEC)
    if resp.status_code != 200:
        raise RuntimeError(f"HTTP {resp.status_code}: {resp.text[:200]}")

    text = resp.text.strip()
    try:
        data = json.loads(text)
    except json.JSONDecodeError as e:
        raise RuntimeError(f"JSON 파싱 실패 (응답 앞부분: {text[:200]}): {e}") from e

    response = data.get("response", data)
    header = response.get("header", {})
    body = response.get("body", {})
    result_code = str(header.get("resultCode", "99"))
    result_msg = str(header.get("resultMsg", "unknown"))

    if result_code not in SUCCESS_RESULT_CODES:
        raise ApiResultError(result_code, result_msg)

    items = normalize_items(body.get("items"))
    total_count = int(body.get("totalCount", len(items)) or 0)
    return {"items": items, "totalCount": total_count}


def call_with_retry(session: requests.Session, operation: str, service_key: str, params: dict[str, str], label: str) -> dict[str, Any]:
    last_err: Exception | None = None
    for attempt in range(MAX_RETRIES + 1):
        try:
            return call_once(session, operation, service_key, params)
        except ApiResultError as e:
            last_err = e
            if e.result_code not in RETRYABLE_RESULT_CODES or attempt == MAX_RETRIES:
                raise
        except Exception as e:  # 네트워크 오류/타임아웃/파싱 실패 등은 재시도
            last_err = e
            if attempt == MAX_RETRIES:
                raise
        delay = RETRY_BASE_DELAY_SEC * (2 ** attempt)
        log(f"  재시도 {attempt + 1}/{MAX_RETRIES} ({label}): {last_err} -> {delay:.1f}s 대기")
        time.sleep(delay)
    raise last_err  # pragma: no cover


def fetch_window(session: requests.Session, business_type: str, operation: str, service_key: str, win_start: date, win_end: date) -> list[dict[str, Any]]:
    all_items: list[dict[str, Any]] = []
    page_no = 1
    while page_no <= MAX_PAGES:
        label = f"{business_type} {win_start}~{win_end} p{page_no}"
        result = call_with_retry(
            session,
            operation,
            service_key,
            {
                "inqryDiv": "1",
                "inqryBgnDt": to_api_datetime(win_start, end_of_day=False),
                "inqryEndDt": to_api_datetime(win_end, end_of_day=True),
                "numOfRows": str(NUM_OF_ROWS),
                "pageNo": str(page_no),
            },
            label,
        )
        items = result["items"]
        all_items.extend(items)
        if len(all_items) >= result["totalCount"] or len(items) < NUM_OF_ROWS or not items:
            break
        page_no += 1
        time.sleep(REQUEST_INTERVAL_SEC)
    return all_items


# ---------------------------------------------------------------------------
# 체크포인트 (진행상황 캐시) - 중간에 끊겨도 이어서 실행 가능하게
# ---------------------------------------------------------------------------


def load_progress() -> set[str]:
    if not PROGRESS_PATH.exists():
        return set()
    return set(json.loads(PROGRESS_PATH.read_text(encoding="utf-8")))


def save_progress(done: set[str]) -> None:
    PROGRESS_PATH.write_text(json.dumps(sorted(done), ensure_ascii=False, indent=2), encoding="utf-8")


def append_raw_items(business_type: str, items: list[dict[str, Any]]) -> None:
    with RAW_CACHE_PATH.open("a", encoding="utf-8") as f:
        for item in items:
            f.write(json.dumps({"business_type": business_type, "item": item}, ensure_ascii=False) + "\n")


def load_cached_items() -> list[tuple[str, dict[str, Any]]]:
    if not RAW_CACHE_PATH.exists():
        return []
    result = []
    with RAW_CACHE_PATH.open(encoding="utf-8") as f:
        for line in f:
            line = line.strip()
            if not line:
                continue
            rec = json.loads(line)
            result.append((rec["business_type"], rec["item"]))
    return result


def run_fetch(service_key: str) -> None:
    CACHE_DIR.mkdir(parents=True, exist_ok=True)
    done = load_progress()
    windows = date_windows(START_DATE, END_DATE, WINDOW_DAYS)
    total_jobs = len(windows) * len(OPERATIONS)
    job_i = 0

    with requests.Session() as session:
        printed_sample = False
        for business_type, operation in OPERATIONS.items():
            for win_start, win_end in windows:
                job_i += 1
                job_key = f"{business_type}:{win_start.isoformat()}:{win_end.isoformat()}"
                if job_key in done:
                    continue

                log(f"[{job_i}/{total_jobs}] {business_type} {win_start} ~ {win_end} 조회 중...")
                try:
                    items = fetch_window(session, business_type, operation, service_key, win_start, win_end)
                except Exception as e:
                    log(f"  실패 (건너뜀, 나중에 재실행하면 이 구간만 다시 시도됩니다): {e}")
                    time.sleep(REQUEST_INTERVAL_SEC)
                    continue

                if items and not printed_sample:
                    printed_sample = True
                    log("  --- 첫 응답 원본 필드명 확인용 샘플 (필드명이 다르면 FIELD_CANDIDATES 수정) ---")
                    log(f"  키 목록: {sorted(items[0].keys())}")

                append_raw_items(business_type, items)
                done.add(job_key)
                save_progress(done)
                log(f"  -> {len(items)}건 수집 (누적 완료 {len(done)}/{total_jobs} 구간)")
                time.sleep(REQUEST_INTERVAL_SEC)

    log("전체 수집 완료.")


# ---------------------------------------------------------------------------
# 필터링
# ---------------------------------------------------------------------------


@dataclass
class FilteredRow:
    notice_no: str
    title: str
    org: str | None
    demand_org: str | None
    contract_method: str | None
    budget: str | None
    notice_date: str | None
    close_date: str | None
    region_limit: str | None
    detail_url: str | None
    matched_keywords: list[str] = field(default_factory=list)
    matched_industry: list[str] = field(default_factory=list)


def evaluate(business_type: str, item: dict[str, Any]) -> FilteredRow | None:
    title = pick(item, FIELD_CANDIDATES["title"]) or ""
    region_limit = pick(item, FIELD_CANDIDATES["region_limit"]) or ""
    industry_text = pick(item, FIELD_CANDIDATES["industry_text"]) or ""

    region_match = REGION_KEYWORD in region_limit
    matched_keywords = [k for k in TITLE_KEYWORDS if k in title]
    matched_industry = [name for name in INDUSTRY_NAMES if name in industry_text]

    keyword_match = len(matched_keywords) > 0
    code_match = len(matched_industry) > 0

    if MATCH_MODE == "OR":
        passed = region_match and (keyword_match or code_match)
    else:  # AND (기본값, 사용자가 명시한 "모두 만족")
        passed = region_match and keyword_match and code_match

    if not passed:
        return None

    notice_no = pick(item, FIELD_CANDIDATES["notice_no"]) or ""
    return FilteredRow(
        notice_no=notice_no,
        title=title,
        org=pick(item, FIELD_CANDIDATES["org"]),
        demand_org=pick(item, FIELD_CANDIDATES["demand_org"]),
        contract_method=pick(item, FIELD_CANDIDATES["contract_method"]),
        budget=pick(item, FIELD_CANDIDATES["budget"]),
        notice_date=pick(item, FIELD_CANDIDATES["notice_date"]),
        close_date=pick(item, FIELD_CANDIDATES["close_date"]),
        region_limit=region_limit or None,
        detail_url=pick(item, FIELD_CANDIDATES["detail_url"]),
        matched_keywords=matched_keywords,
        matched_industry=matched_industry,
    )


def print_funnel_stats(records: list[tuple[str, dict[str, Any]]]) -> None:
    """AND/OR 중 어느 쪽이 나을지 판단할 수 있도록 조건별 건수를 보여준다."""
    region_only = keyword_only_after_region = code_only_after_region = both_after_region = 0
    for _bt, item in records:
        title = pick(item, FIELD_CANDIDATES["title"]) or ""
        region_limit = pick(item, FIELD_CANDIDATES["region_limit"]) or ""
        industry_text = pick(item, FIELD_CANDIDATES["industry_text"]) or ""

        if REGION_KEYWORD not in region_limit:
            continue
        region_only += 1
        k = any(kw in title for kw in TITLE_KEYWORDS)
        c = any(name in industry_text for name in INDUSTRY_NAMES)
        if k:
            keyword_only_after_region += 1
        if c:
            code_only_after_region += 1
        if k and c:
            both_after_region += 1

    log("--- 조건별 건수 (전체 수집 건 대비, MATCH_MODE 판단용) ---")
    log(f"  강원 지역제한 포함: {region_only}건")
    log(f"    ㄴ 그 중 제목 키워드까지 매칭: {keyword_only_after_region}건")
    log(f"    ㄴ 그 중 업종코드까지 매칭:   {code_only_after_region}건")
    log(f"    ㄴ 키워드 AND 업종코드 모두:  {both_after_region}건  (현재 MATCH_MODE=AND 결과와 동일)")
    log(f"    ㄴ 키워드 OR 업종코드:        {keyword_only_after_region + code_only_after_region - both_after_region}건  (MATCH_MODE=OR 이면 이 값)")


# ---------------------------------------------------------------------------
# 엑셀 출력
# ---------------------------------------------------------------------------

COLUMNS = [
    ("공고번호", "notice_no", 16),
    ("공고명", "title", 46),
    ("발주기관", "org", 20),
    ("수요기관", "demand_org", 20),
    ("계약방법", "contract_method", 16),
    ("추정가격(예산)", "budget", 16),
    ("공고일", "notice_date", 12),
    ("입찰마감일", "close_date", 14),
    ("지역제한", "region_limit", 18),
    ("원문 URL", "detail_url", 40),
]


def write_excel(rows: list[FilteredRow], out_path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "강원_전시체험_공고"

    header_font = Font(name="Arial", bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    body_font = Font(name="Arial")

    for col_i, (header, _key, width) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_i, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.column_dimensions[get_column_letter(col_i)].width = width
    ws.freeze_panes = "A2"

    for row_i, row in enumerate(rows, start=2):
        for col_i, (_header, key, _width) in enumerate(COLUMNS, start=1):
            value: Any = getattr(row, key)
            if key == "budget" and value is not None:
                try:
                    value = int(str(value).replace(",", ""))
                except ValueError:
                    pass
            cell = ws.cell(row=row_i, column=col_i, value=value)
            cell.font = body_font
            if key == "budget" and isinstance(value, int):
                cell.number_format = "#,##0"
            if key == "detail_url" and value:
                cell.hyperlink = value
                cell.font = Font(name="Arial", color="0563C1", underline="single")

    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)


# ---------------------------------------------------------------------------
# 메인
# ---------------------------------------------------------------------------


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--skip-fetch", action="store_true", help="API 재호출 없이 캐시된 데이터로만 필터/엑셀 생성")
    args = parser.parse_args()

    service_key = os.environ.get("NARA_SERVICE_KEY", "").strip()

    if not args.skip_fetch:
        if not service_key:
            print("환경변수 NARA_SERVICE_KEY 가 설정되지 않았습니다.", file=sys.stderr)
            print('예: export NARA_SERVICE_KEY="발급받은 Decoding 인증키"', file=sys.stderr)
            return 1
        run_fetch(service_key)
    else:
        log("--skip-fetch: API 호출 없이 캐시된 원본 데이터만 사용합니다.")

    records = load_cached_items()
    log(f"캐시에서 원본 {len(records)}건 로드 (용역+공사, 15일 창 전체 합산, 중복 제거 전)")

    # 공고번호 기준 중복 제거 (창 경계 등으로 중복 수집될 수 있음)
    seen: set[str] = set()
    deduped: list[tuple[str, dict[str, Any]]] = []
    for bt, item in records:
        no = pick(item, FIELD_CANDIDATES["notice_no"]) or json.dumps(item, ensure_ascii=False)
        if no in seen:
            continue
        seen.add(no)
        deduped.append((bt, item))
    log(f"중복 제거 후 {len(deduped)}건")

    print_funnel_stats(deduped)

    rows: list[FilteredRow] = []
    for bt, item in deduped:
        row = evaluate(bt, item)
        if row:
            rows.append(row)

    log(f"필터 조건(MATCH_MODE={MATCH_MODE}) 통과: 총 {len(rows)}건")

    out_path = OUTPUT_DIR / OUTPUT_FILENAME
    write_excel(rows, out_path)
    log(f"엑셀 저장 완료: {out_path}")
    log(f"=== 총 {len(rows)}건 ===")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
